#!/usr/bin/env python
"""Eval-harness for risikovurderingsmotoren (briefens §12).

Kører hele pipelinen (analyze → generate → render → verify) mod kildedokumenter
og sammenligner med en manuelt udfyldt ground-truth-vurdering på 9 dimensioner.

Brug:
    python scripts/eval_risikovurdering.py \
        --source "sti/til/kildedok.docx" [--source flere.pdf ...] \
        --systemnavn Voicecraft \
        --ground-truth "sti/til/Databeskyttelsesretlig risikovurdering - Voicecraft.docx" \
        [--answers answers.json] [--out /tmp/eval-ud] [--json]

Dimensioner (PASS/FAIL + detalje):
  1. struktur_tabeller    — genereret docx har skabelonens 8 tabeller
  2. risici_antal         — 8-12 risici
  3. verify_valid         — verifier finder ingen placeholder-rester/mangler
  4. felttekster_udfyldt  — alle 14 prosa-felter udfyldt
  5. leverandoer_match    — ekstraheret leverandør optræder i ground-truth
  6. hosting_match        — hosting-lokation genfindes i ground-truth
  7. kategorier_match     — persondata-kategorier konsistente med ground-truth
  8. msa_flags_genfundet  — §-referencer fra ground-truth fanges af motoren
  9. risiko_tema_overlap  — tematisk overlap mellem genererede + manuelle risici

Exit-kode: 0 = alle obligatoriske dimensioner PASS, 1 = mindst én FAIL, 2 = fejl.
LLM-tid: ~2-3 min pr. system på lokal model.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import time

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(_ROOT, ".env"))
except ImportError:
    pass


def _stopword_tokens(text: str) -> set[str]:
    """Indholds-tokens (>4 tegn, lowercased) til tema-overlap."""
    return {
        t for t in re.findall(r"[a-zæøåA-ZÆØÅ]{5,}", text.lower())
        if t not in {
            "kommunens", "kommunen", "kalundborg", "systemet", "system",
            "risiko", "risici", "behandling", "personoplysninger", "mellem",
            "herunder", "eller", "denne", "dette", "deres", "kunne", "skulle",
        }
    }


def run_eval(args) -> dict:
    from docx import Document

    from src.services.risk_assessment import orchestrator
    from src.services.risk_assessment.docx_assembler import assemble_docx, output_filename
    from src.services.risk_assessment.document_extract import extract_document
    from src.services.risk_assessment.verifier import verify_docx

    t0 = time.time()

    # ---- Kør pipeline -----------------------------------------------------
    files = []
    for path in args.source:
        with open(path, "rb") as f:
            files.append((os.path.basename(path), f.read()))

    print(f"[1/4] analyze ({len(files)} dokument(er))…", flush=True)
    ar = orchestrator.analyze(files, systemnavn=args.systemnavn, timeout=240)
    facts = ar.facts

    answers = {"scope": "Begge", "tilgang": "Idealiseret (best practices på plads)"}
    if args.answers:
        with open(args.answers) as f:
            answers.update(json.load(f))

    print(f"[2/4] generate (LLM — kan tage 2-3 min)…", flush=True)
    rv = orchestrator.generate(facts, answers, timeout=300)

    print(f"[3/4] render + verify…", flush=True)
    docx_bytes = assemble_docx(rv)
    ver = verify_docx(docx_bytes, expected_n_risks=len(rv.risici), facts=rv.facts)

    out_dir = args.out or "/tmp"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, output_filename(args.systemnavn))
    with open(out_path, "wb") as f:
        f.write(docx_bytes)

    # ---- Ground-truth -----------------------------------------------------
    print(f"[4/4] sammenligner med ground-truth…", flush=True)
    with open(args.ground_truth, "rb") as f:
        gt_bytes = f.read()
    gt_text = extract_document(os.path.basename(args.ground_truth), gt_bytes).text
    gt_lower = gt_text.lower()
    gt_doc = Document(io.BytesIO(gt_bytes))

    gen_doc = Document(io.BytesIO(docx_bytes))

    dims: dict[str, dict] = {}

    def dim(name, ok, detail, *, mandatory=True):
        dims[name] = {"pass": bool(ok), "detail": detail, "mandatory": mandatory}

    # 1. Struktur
    dim("struktur_tabeller", len(gen_doc.tables) == 8,
        f"genereret har {len(gen_doc.tables)} tabeller (forventet 8)")

    # 2. Risici-antal
    n = len(rv.risici)
    dim("risici_antal", 8 <= n <= 12, f"{n} risici (krav 8-12)")

    # 3. Verifier
    dim("verify_valid", ver.valid,
        "ingen problemer" if ver.valid else "; ".join(ver.problems[:4]))

    # 4. Felttekster
    tomme = [k for k, v in rv.alle_felttekster().items() if not v.strip()]
    dim("felttekster_udfyldt", not tomme,
        "14/14 udfyldt" if not tomme else f"tomme: {', '.join(tomme)}")

    # 5. Leverandør genfindes i ground-truth
    lev = (facts.leverandoer_navn or "").strip()
    if facts.internt_udviklet:
        dim("leverandoer_match", True, "internt udviklet — ingen leverandør forventet")
    elif lev:
        # match på første ord (fx 'Velatir' i 'Velatir ApS')
        first = lev.split()[0].lower()
        dim("leverandoer_match", first in gt_lower, f"{lev!r} {'genfundet' if first in gt_lower else 'IKKE fundet'} i ground-truth")
    else:
        dim("leverandoer_match", False, "ingen leverandør ekstraheret", mandatory=False)

    # 6. Hosting genfindes
    host_tokens = _stopword_tokens(facts.hosting_lokation)
    host_hit = any(t in gt_lower for t in host_tokens) if host_tokens else False
    dim("hosting_match", host_hit,
        f"hosting={facts.hosting_lokation!r} → {'genfundet' if host_hit else 'ikke fundet'}",
        mandatory=False)

    # 7. Kategorier konsistente (ekstraherede kategorier nævnes i ground-truth)
    kats = [k.value for k in facts.persondata_kategorier]
    kat_keywords = {"almindelige": "almindelige", "følsomme": "følsom", "cpr": "cpr", "strafbare": "strafbare", "ingen": None}
    misses = [k for k in kats if kat_keywords.get(k) and kat_keywords[k] not in gt_lower]
    dim("kategorier_match", not misses,
        f"kategorier {kats} — {'alle genfundet' if not misses else 'ikke i GT: ' + str(misses)}",
        mandatory=False)

    # 8. MSA §-flag fra ground-truth fanges af motoren (kun hvis GT har §-refs)
    gt_pars = set(re.findall(r"§\s*(\d+(?:\.\d+)?)", gt_text))
    gen_flag_text = " ".join(facts.msa_risiko_klausuler) + " " + " ".join(r.risiko for r in rv.risici)
    gen_pars = set(re.findall(r"§\s*(\d+(?:\.\d+)?)", gen_flag_text))
    if gt_pars:
        found = gt_pars & gen_pars
        frac = len(found) / len(gt_pars)
        dim("msa_flags_genfundet", frac >= 0.3,
            f"{len(found)}/{len(gt_pars)} §-refs genfundet ({sorted(found)})", mandatory=False)
    else:
        dim("msa_flags_genfundet", True, "ground-truth har ingen §-referencer", mandatory=False)

    # 9. Tematisk overlap mellem genererede risici og GT-risikoskema (Table 5)
    gt_risk_text = ""
    if len(gt_doc.tables) > 5:
        for row in gt_doc.tables[5].rows[1:]:
            gt_risk_text += " " + row.cells[0].text
    gt_tokens = _stopword_tokens(gt_risk_text)
    gen_tokens = _stopword_tokens(" ".join(r.risiko + " " + r.hvorfor for r in rv.risici))
    overlap = len(gt_tokens & gen_tokens) / max(1, len(gt_tokens))
    dim("risiko_tema_overlap", overlap >= 0.20,
        f"{overlap:.0%} token-overlap med ground-truth-risici (krav ≥20%)", mandatory=False)

    elapsed = time.time() - t0
    mandatory_fail = [k for k, v in dims.items() if v["mandatory"] and not v["pass"]]
    return {
        "systemnavn": args.systemnavn,
        "elapsed_s": round(elapsed, 1),
        "n_risici": n,
        "output_docx": out_path,
        "dimensions": dims,
        "all_mandatory_pass": not mandatory_fail,
        "mandatory_failures": mandatory_fail,
    }


def main() -> int:
    p = argparse.ArgumentParser(description="Eval af risikovurderingsmotoren mod ground-truth")
    p.add_argument("--source", action="append", required=True, help="Kildedokument (gentag for flere)")
    p.add_argument("--systemnavn", required=True)
    p.add_argument("--ground-truth", required=True, help="Manuelt udfyldt vurdering (docx)")
    p.add_argument("--answers", help="JSON-fil med svar på afklarende spørgsmål")
    p.add_argument("--out", help="Output-mappe til genereret docx (default /tmp)")
    p.add_argument("--json", action="store_true", help="Maskinlæsbart JSON-output")
    args = p.parse_args()

    try:
        result = run_eval(args)
    except Exception as exc:
        print(f"FATAL: {exc}", file=sys.stderr)
        return 2

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(f"\n=== EVAL: {result['systemnavn']} ({result['elapsed_s']}s, {result['n_risici']} risici) ===")
        for name, d in result["dimensions"].items():
            mark = "✓" if d["pass"] else ("✗" if d["mandatory"] else "⚠")
            print(f"  {mark} {name:24s} {d['detail']}")
        print(f"\nDocx: {result['output_docx']}")
        print("RESULTAT:", "PASS" if result["all_mandatory_pass"] else f"FAIL ({', '.join(result['mandatory_failures'])})")

    return 0 if result["all_mandatory_pass"] else 1


if __name__ == "__main__":
    sys.exit(main())
