"""Deterministisk DOCX-udfyldning af Kalundborg Kommunes risikovurderings-skabelon.

Porteret fra skill-risikovurdering/scripts/fill_template.py. Nøgle-teknik:
felter der skal udfyldes er markeret med GUL highlight i master-templaten
(`run.font.highlight_color is not None`). Vi erstatter den gule run-tekst og
rydder bracket-markører i nabo-runs.

Master-templatens struktur (8 tabeller):
  para 3        : systemnavn (gul "Plan2learn"-placeholder)
  Table 0       : formål (P1), omfang (P3), ansvarlige (P5)
  Table 1       : baggrund (P2), funktionalitet (P5), interessenter (P7)
  Table 2       : person (P0), lokation (P1), adgang (P2), sårbarheder (P3)
  Table 3, 4    : "Eksempler" guidance — RØR IKKE
  Table 5       : risikoskema, 16 rækker (header + 1 template + 14 tomme)
  Table 6       : tiltag (yp[0]), ansvarlige_tiltag (yp[2]); yp[1] er Eksempler-label
  Table 7       : kontrol (yp[0]), opdatering (yp[1]) — label står foran som ikke-gul tekst

Master-filen ændres ALDRIG — vi læser fra den og skriver til en BytesIO/ny fil.
"""

import io
import logging
import os
from copy import deepcopy

from docx import Document

from src.services.risk_assessment.models import Risikovurdering

logger = logging.getLogger("bifrost.risk_assessment.assembler")

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# Default sti til master-template i repo'et
_DEFAULT_TEMPLATE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))),
    "data", "templates", "risikovurdering_master.docx",
)


def get_template_path() -> str:
    """Returnér sti til master-template. Override via RISK_TEMPLATE_PATH env."""
    return os.getenv("RISK_TEMPLATE_PATH", _DEFAULT_TEMPLATE)


# ---- gul-highlight helpers (porteret fra fill_template.py) ----------------


_PUNCT_ONLY = set(".,;: \t")
_SENTENCE_END = (".", "!", "?", ":")


def _replace_yellow_in_paragraph(p, new_text: str) -> bool:
    """Erstat gul-markeret tekst i en paragraf og ryd bracket-scaffolding.

    Master-templaten omkranser gule felter med `[ ... ]`-markører fordelt på
    ikke-gule nabo-runs. Markørerne ligger ikke altid direkte op ad den gule
    region (der kan være en mellemliggende `.`-run), så vi scanner i stedet for
    kun at tjekke den umiddelbare nabo.

    Regler:
      - Leading: fjern ét afsluttende `[` fra runnen før første gule (bevarer
        evt. label som "Kontrolmekanismer: ").
      - Trailing: i runs efter sidste gule (indtil næste gule/slut) fjernes
        `[`/`]`-tegn. En run der derefter kun består af tegnsætning ryddes KUN
        hvis new_text allerede ender på sætningstegn (undgår dobbelt-punktum;
        bevarer dog et punktum hvis new_text mangler et).
    """
    runs = p.runs
    yellow_idx = [i for i, r in enumerate(runs) if r.font.highlight_color is not None]
    if not yellow_idx:
        return False
    first, last = yellow_idx[0], yellow_idx[-1]

    new_text = (new_text or "").strip()
    runs[first].text = new_text
    for i in yellow_idx[1:]:
        runs[i].text = ""

    # Leading: strip ét trailing '[' fra runnen lige før første gule
    if first > 0:
        prev = runs[first - 1]
        if prev.font.highlight_color is None and "[" in prev.text:
            idx = prev.text.rfind("[")
            prev.text = (prev.text[:idx] + prev.text[idx + 1:]).rstrip()
            if prev.text and not prev.text.endswith((" ", ":")):
                prev.text += " "

    ends_with_punct = new_text.endswith(_SENTENCE_END)

    # Trailing: scan ALLE runs efter sidste gule indtil næste gule / slut.
    # Fjern bracket-tegn fra hver; ryd redundante tegnsætnings-runs men bevar
    # ét punktum hvis new_text mangler sætningsafslutning.
    preserved_period = False
    for j in range(last + 1, len(runs)):
        r = runs[j]
        if r.font.highlight_color is not None:
            break
        if "[" in r.text or "]" in r.text:
            r.text = r.text.replace("[", "").replace("]", "")
        stripped = r.text.strip()
        if not stripped:
            r.text = ""
            continue
        if all(ch in _PUNCT_ONLY for ch in stripped):
            if ends_with_punct or preserved_period:
                r.text = ""
            else:
                preserved_period = True  # behold denne som sætningsafslutning
        # ellers: rigtig tekst (sjældent) — lad stå
    return True


def _yellow_paragraphs(cell):
    return [
        p for p in cell.paragraphs
        if any(r.font.highlight_color is not None for r in p.runs)
    ]


def _fill_yellow_paragraphs(cell, texts: list):
    """Fyld de gule paragraffer i en celle i rækkefølge (None = spring over)."""
    yps = _yellow_paragraphs(cell)
    for i, txt in enumerate(texts):
        if i < len(yps) and txt is not None and str(txt).strip():
            _replace_yellow_in_paragraph(yps[i], str(txt))


def _fill_risk_row(row, risk):
    cells = row.cells
    # C0: risiko
    yps = _yellow_paragraphs(cells[0])
    if yps:
        _replace_yellow_in_paragraph(yps[0], risk.risiko)
    # C1: konsekvens
    yps = _yellow_paragraphs(cells[1])
    if yps:
        _replace_yellow_in_paragraph(yps[0], risk.konsekvens_beskrivelse)
    # C2: sandsynlighed
    yps = _yellow_paragraphs(cells[2])
    if yps:
        _replace_yellow_in_paragraph(yps[0], risk.sandsynlighed.value)
    # C3: score (yp[0]) + hvorfor (yp[1])
    yps = _yellow_paragraphs(cells[3])
    if len(yps) >= 2:
        _replace_yellow_in_paragraph(yps[0], risk.score.value)
        _replace_yellow_in_paragraph(yps[1], risk.hvorfor or "")
    elif yps:
        combined = f"{risk.score.value} – {risk.hvorfor}" if risk.hvorfor else risk.score.value
        _replace_yellow_in_paragraph(yps[0], combined)


# ---- hovedflow ------------------------------------------------------------


def assemble_docx(rv: Risikovurdering, *, template_path: str | None = None) -> bytes:
    """Byg et udfyldt risikovurderings-dokument og returnér som bytes.

    Læser master-templaten, udfylder alle gule felter + risikoskemaet, og
    returnerer DOCX som bytes (klar til HTTP-download eller fil-skrivning).
    """
    src = template_path or get_template_path()
    if not os.path.exists(src):
        raise FileNotFoundError(f"Master-template ikke fundet: {src}")

    doc = Document(src)

    # Indledning para 3 — erstat gul "Plan2learn" med systemnavn
    systemnavn = rv.facts.systemnavn or "systemet"
    for r in doc.paragraphs[3].runs:
        if r.font.highlight_color is not None and "Plan2learn" in r.text:
            r.text = systemnavn
    # Titel para 0 — opdatér systemnavn efter bindestreg hvis present
    _update_title(doc, systemnavn)

    # Table 0: Formål, Omfang, Ansvarlige
    _fill_yellow_paragraphs(
        doc.tables[0].rows[0].cells[0],
        [rv.formaal_tekst, rv.omfang_tekst, rv.ansvarlige_tekst],
    )
    # Table 1: Baggrund, Funktionalitet, Interessenter
    _fill_yellow_paragraphs(
        doc.tables[1].rows[0].cells[0],
        [rv.baggrund_tekst, rv.funktionalitet_tekst, rv.interessenter_tekst],
    )
    # Table 2: Person, Lokation, Adgang, Sårbarheder
    _fill_yellow_paragraphs(
        doc.tables[2].rows[0].cells[0],
        [
            rv.personoplysninger_tekst,
            rv.lokationer_tekst,
            rv.adgangsrettigheder_tekst,
            rv.saarbarheder_tekst,
        ],
    )

    # Table 5: Risikoskema — fjern tomme rækker, duplikér template-rækken
    _fill_risk_table(doc.tables[5], rv.risici)

    # Table 6: Tiltag (yp[0]) + Ansvarlige (yp[2]); yp[1] = Eksempler-label
    cell6 = doc.tables[6].rows[0].cells[0]
    yps6 = _yellow_paragraphs(cell6)
    if len(yps6) >= 1 and rv.tiltag_tekst.strip():
        _replace_yellow_in_paragraph(yps6[0], rv.tiltag_tekst)
    if len(yps6) >= 3 and rv.ansvarlige_tiltag_tekst.strip():
        _replace_yellow_in_paragraph(yps6[2], rv.ansvarlige_tiltag_tekst)

    # Table 7: Kontrolmekanismer + Opdatering (labels står foran som ikke-gul tekst)
    _fill_yellow_paragraphs(
        doc.tables[7].rows[0].cells[0],
        [rv.kontrolmekanismer_tekst, rv.opdatering_tekst],
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _update_title(doc, systemnavn: str):
    """Opdatér titlen i para 0 så den ender med systemnavnet."""
    p = doc.paragraphs[0]
    base = "Skabelon til databeskyttelsesretlig risikovurdering"
    full = f"{base} - {systemnavn}"
    # Hvis hele titlen ligger i én run, sæt den; ellers sæt første run + ryd resten
    if p.runs:
        p.runs[0].text = full
        for r in p.runs[1:]:
            r.text = ""


def _fill_risk_table(table, risici):
    """Behold header (række 0), brug række 1 som template, fyld N risici.

    Master har 16 rækker (header + template + 14 tomme). Vi fjerner alt efter
    række 1, fylder template-rækken med risk[0], og duplikerer den for resten.
    """
    if not risici:
        logger.warning("Ingen risici at indsætte — risikoskema forbliver tomt")
        return

    tbl = table._tbl
    trs = tbl.findall(f"{{{_W_NS}}}tr")
    # Fjern alle rækker efter den første template-række (index 2+)
    for tr in trs[2:]:
        tbl.remove(tr)

    template_row = table.rows[1]
    _fill_risk_row(template_row, risici[0])

    template_tr = template_row._tr
    parent = template_tr.getparent()
    idx = list(parent).index(template_tr)
    for _ in risici[1:]:
        new_tr = deepcopy(template_tr)
        parent.insert(idx + 1, new_tr)
        idx += 1
    # Re-fyld de duplikerede rækker (deepcopy kopierede risk[0]'s tekst)
    for i, risk in enumerate(risici[1:], start=2):
        _fill_risk_row(table.rows[i], risk)


def assemble_to_file(rv: Risikovurdering, output_path: str, *, template_path: str | None = None) -> str:
    """Som assemble_docx men skriver til fil. Returnér output_path."""
    data = assemble_docx(rv, template_path=template_path)
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    with open(output_path, "wb") as f:
        f.write(data)
    return output_path


def output_filename(systemnavn: str) -> str:
    """Standard-filnavn jf. Kalundborg-konvention.

    Sanitizer alle tegn der er ulovlige i Windows/macOS-filnavne — ikke kun '/'
    (Windows afviser fx 'System*V2.docx' ved download).
    """
    import re as _re
    safe = _re.sub(r'[\\/:*?"<>|]', "-", (systemnavn or "system")).strip()
    safe = _re.sub(r"-{2,}", "-", safe).strip("- ") or "system"
    return f"Databeskyttelsesretlig risikovurdering - {safe}.docx"
