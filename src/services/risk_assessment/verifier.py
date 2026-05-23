"""Verifikation af et færdigt risikovurderings-dokument.

Porteret fra skill-risikovurdering/scripts/verify.py + udvidet med Kalundborg-
compliance-tjek (jf. Retningslinjer for IT-anskaffelser + AI-tjekliste).

Tjekker:
  - "Plan2learn"-placeholder er erstattet i indledningen
  - Ingen gule placeholder-fraser er tilbage (uerstattede felter)
  - Risikoskema (Table 5) har 1 header + N risici (ikke 16 tomme rækker)
  - Table 7 har ikke dobbelt label
  - Ingen rester af tidligere systemnavne
  - (valgfrit, hvis facts gives) Kommunale compliance-issues:
    * Udbudspligt-mismatch nævnt i tekst hvis facts.udbudspligt_mismatch()
    * Mindst én kommunal aktør (CIO/DPO) nævnt i ansvarlige-tekst
    * Fagområde + særlovgivning nævnt hvis facts.fagomraade er sat

Returnerer en VerifyResult med problems-liste (tom = OK).
"""

import io
import logging
from dataclasses import dataclass, field
from typing import Optional, TYPE_CHECKING

from docx import Document

if TYPE_CHECKING:
    from src.services.risk_assessment.models import SystemFacts

logger = logging.getLogger("bifrost.risk_assessment.verifier")


# Placeholder-fraser fra master-templaten — hvis disse er tilbage, blev feltet ikke fyldt
_PLACEHOLDER_PHRASES = [
    "Beskriv formålet med risikovurderingen",
    "Beskriv kort projektets formål",
    "Beskriv de tekniske og organisatoriske",
    "Beskriv hvilke personoplysninger",
    "Beskriv, hvordan risici",
    "Beskriv, hvordan risici løbende overvåges",
    "Angiv, hvilke dele af projektet",
    "Angiv de vigtigste funktioner",
    "Angiv, hvem der er ansvarlig",
    "Angiv, hvornår og hvordan risikovurderingen",
    "Identificer de ansvarlige",
    "Identificer relevante interessenter",
    "Overvej sikkerheden omkring",
    "Overvej styringen af adgangsrettigheder",
    "Overvej om der er nogen relevante sårbarheder",
    "indsæt risici",
    "indsæt konsekvensen",
    "Lav, middel eller høj",
]

# Tidligere systemnavne der ikke må lække ind i et nyt dokument
_FORBIDDEN_LEFTOVERS = ["Plan2learn", "[INDSÆT", "TODO", "<SYSTEMNAVN>"]


@dataclass
class VerifyResult:
    valid: bool
    n_risk_rows: int = 0
    problems: list = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "valid": self.valid,
            "n_risk_rows": self.n_risk_rows,
            "problems": self.problems,
        }


def verify_docx(
    data: bytes,
    *,
    expected_n_risks: Optional[int] = None,
    systemnavn: Optional[str] = None,
    facts: Optional["SystemFacts"] = None,
) -> VerifyResult:
    """Verificér et færdigt dokument (bytes). Returnér VerifyResult.

    Hvis facts gives, køres også Kalundborg-specifikke compliance-tjek:
    udbudspligt-mismatch, kommunale aktører nævnt, særlovgivning refereret.
    """
    doc = Document(io.BytesIO(data))
    problems: list[str] = []

    # 1. Indledning — "Plan2learn" må ikke optræde
    if len(doc.paragraphs) > 3 and "Plan2learn" in doc.paragraphs[3].text:
        problems.append("Para 3 indeholder stadig placeholder 'Plan2learn'")

    # 2. Gule placeholder-fraser i indholdstabeller (0,1,2,5,6,7)
    for ti in [0, 1, 2, 5, 6, 7]:
        if ti >= len(doc.tables):
            continue
        t = doc.tables[ti]
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    for r in p.runs:
                        if r.font.highlight_color is None:
                            continue
                        txt = r.text.strip()
                        if not txt:
                            continue
                        for phrase in _PLACEHOLDER_PHRASES:
                            if phrase in txt:
                                problems.append(
                                    f"Table {ti} R{ri}C{ci}: placeholder ikke erstattet: {txt[:60]!r}"
                                )
                                break

    # 3. Risikoskema antal rækker
    n_rows = len(doc.tables[5].rows) if len(doc.tables) > 5 else 0
    n_risks = max(0, n_rows - 1)
    if n_rows > 13:
        problems.append(f"Table 5 har {n_rows} rækker — tomme rækker bør slettes")
    elif n_rows < 3:
        problems.append(f"Table 5 har kun {n_rows} rækker — for få risici")
    if expected_n_risks is not None and n_risks != expected_n_risks:
        problems.append(
            f"Risikoskema har {n_risks} risici, forventet {expected_n_risks}"
        )

    # 4. Table 7 dobbelt label
    if len(doc.tables) > 7:
        t7 = doc.tables[7].rows[0].cells[0].text
        if t7.count("Kontrolmekanismer:") > 1:
            problems.append("Table 7: 'Kontrolmekanismer:' optræder mere end én gang")
        if t7.count("Opdatering af risikovurdering:") > 1:
            problems.append("Table 7: 'Opdatering af risikovurdering:' optræder mere end én gang")

    # 5. Forbudte rester (placeholder + tidligere systemnavne)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text
    for forbudt in _FORBIDDEN_LEFTOVERS:
        if forbudt in full_text:
            problems.append(f"Rester af placeholder/tidligere system: '{forbudt}'")

    # 6. Kalundborg-compliance-tjek (kun hvis facts er givet)
    if facts is not None:
        full_lower = full_text.lower()

        # 6a. Udbudspligt-mismatch — hvis fakta indikerer det, skal teksten nævne det
        if hasattr(facts, "udbudspligt_mismatch") and facts.udbudspligt_mismatch():
            if not any(t in full_lower for t in ("eu-udbud", "udbudspligt", "udbudslov")):
                problems.append(
                    "Udbudspligt-mismatch i fakta (værdi over tærskel, ikke EU-udbud), "
                    "men teksten nævner ikke EU-udbud/udbudspligt — bør indgå som blocker"
                )

        # 6b. Mindst én kommunal aktør i ansvarlige-tekst (CIO eller DPO)
        kommunale_aktorer = (
            "digitaliserings- og it-chefen", "it-chefen", "dpo",
            "databeskyttelsesrådgiver", "ai-gruppen",
        )
        if len(doc.tables) > 0:
            t0_text = doc.tables[0].rows[0].cells[0].text.lower()
            if not any(a in t0_text for a in kommunale_aktorer):
                problems.append(
                    "Table 0 (ansvarlige) nævner ingen kommunale aktører "
                    "(CIO, DPO, AI-gruppen) — kommunal kontekst mangler"
                )

        # 6c. Hvis fagområde er angivet, skal teksten nævne særlovgivning eller fagområdet
        if getattr(facts, "fagomraade", ""):
            fag = facts.fagomraade.lower()
            saerlov = [s.lower() for s in getattr(facts, "saerlovgivning", [])]
            mentions_fag = fag in full_lower
            mentions_saerlov = any(s in full_lower for s in saerlov) if saerlov else False
            mentions_saerlov_keyword = "særlov" in full_lower or "særlovgivning" in full_lower
            if not (mentions_fag or mentions_saerlov or mentions_saerlov_keyword):
                problems.append(
                    f"Fagområde '{facts.fagomraade}' angivet i fakta, men hverken "
                    f"fagområdet eller særlovgivning er nævnt i teksten"
                )

    return VerifyResult(valid=not problems, n_risk_rows=n_risks, problems=problems)
