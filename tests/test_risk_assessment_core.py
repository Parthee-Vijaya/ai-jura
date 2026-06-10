"""Tests for risikovurderingsmotorens deterministiske kerne:
  - models: normalisering, helpers
  - document_extract: pdf/docx/txt + afvisning
  - docx_assembler: udfyldning mod den RIGTIGE master-template
  - verifier: detekterer placeholder-rester + forkert rækkeantal

LLM-services testes separat med mocks (test_risk_assessment_llm.py).
"""

import io
import os

import pytest
from docx import Document

from src.services.risk_assessment.models import (
    Anskaffelsesvej,
    DataKategori,
    Niveau,
    Risiko,
    Risikovurdering,
    SystemFacts,
    UDBUDSTERSKEL_KR_4AAR,
    normalize_niveau,
)
from src.services.risk_assessment.document_extract import (
    DocumentExtractError,
    extract_document,
    combine_for_prompt,
)
from src.services.risk_assessment.docx_assembler import (
    assemble_docx,
    get_template_path,
    output_filename,
)
from src.services.risk_assessment.verifier import verify_docx


# ---- models --------------------------------------------------------------


class TestModels:
    def test_normalize_niveau_aliases(self):
        assert normalize_niveau("mellem") == Niveau.MIDDEL
        assert normalize_niveau("high") == Niveau.HOEJ
        assert normalize_niveau("Lav-middel") == Niveau.LAV_MIDDEL
        assert normalize_niveau("hoj") == Niveau.HOEJ

    def test_normalize_niveau_fallback(self):
        assert normalize_niveau("garbage") == Niveau.MIDDEL
        assert normalize_niveau("") == Niveau.MIDDEL
        assert normalize_niveau(None) == Niveau.MIDDEL

    def test_har_foelsomme_data(self):
        f = SystemFacts(persondata_kategorier=[DataKategori.CPR])
        assert f.har_foelsomme_data()
        f2 = SystemFacts(persondata_kategorier=[DataKategori.ALMINDELIGE])
        assert not f2.har_foelsomme_data()

    def test_er_ekstern_leverandoer(self):
        ekstern = SystemFacts(leverandoer_navn="Acme", internt_udviklet=False)
        assert ekstern.er_ekstern_leverandoer()
        intern = SystemFacts(leverandoer_navn="", internt_udviklet=True)
        assert not intern.er_ekstern_leverandoer()

    def test_alle_felttekster_count(self):
        rv = Risikovurdering(facts=SystemFacts())
        assert len(rv.alle_felttekster()) == 14


class TestKalundborgFelter:
    """Etape 1+2: nye kommunale procesfelter + helpers."""

    def test_udbudsterskel_konstant(self):
        assert UDBUDSTERSKEL_KR_4AAR == 1_601_944

    def test_er_over_udbudsterskel(self):
        assert SystemFacts(kontraktvaerdi_4aar_kr=2_000_000).er_over_udbudsterskel() is True
        assert SystemFacts(kontraktvaerdi_4aar_kr=500_000).er_over_udbudsterskel() is False
        assert SystemFacts().er_over_udbudsterskel() is None  # ukendt

    def test_udbudspligt_mismatch_over_taerskel_uden_eu_udbud(self):
        f = SystemFacts(
            kontraktvaerdi_4aar_kr=3_000_000,
            anskaffelsesvej=Anskaffelsesvej.SKI_DIREKTE,
        )
        assert f.udbudspligt_mismatch() is True

    def test_udbudspligt_mismatch_eu_udbud_ok(self):
        f = SystemFacts(
            kontraktvaerdi_4aar_kr=3_000_000,
            anskaffelsesvej=Anskaffelsesvej.EU_UDBUD,
        )
        assert f.udbudspligt_mismatch() is False

    def test_udbudspligt_mismatch_under_taerskel(self):
        # Under tærskel = ingen mismatch uanset anskaffelsesvej
        f = SystemFacts(
            kontraktvaerdi_4aar_kr=500_000,
            anskaffelsesvej=Anskaffelsesvej.UNDER_TAERSKEL,
        )
        assert f.udbudspligt_mismatch() is False

    def test_udbudspligt_mismatch_ukendt_ikke_mismatch(self):
        # UKENDT-anskaffelsesvej tæller ikke som mismatch (afventer afklaring)
        f = SystemFacts(
            kontraktvaerdi_4aar_kr=3_000_000,
            anskaffelsesvej=Anskaffelsesvej.UKENDT,
        )
        assert f.udbudspligt_mismatch() is False

    def test_proces_status_count(self):
        assert SystemFacts().proces_status_count() == (0, 9)
        f = SystemFacts(
            dit_involveret_tidligt=True,
            cio_har_underskrevet=True,
            databehandleraftale_indgaaet=True,
        )
        assert f.proces_status_count() == (3, 9)

    def test_anskaffelsesvej_enum_has_6_options(self):
        # SKI direkte/mini, under tærskel, EU-udbud, bygge-anlæg, ukendt
        assert len(list(Anskaffelsesvej)) == 6


class TestClarifyingNewQuestions:
    """Etape 2: nye spørgsmål + apply_answers parsing."""

    def test_build_questions_includes_kalundborg(self):
        from src.services.risk_assessment.clarifying import build_questions
        qs = build_questions(SystemFacts(hosting_lokation="Azure"))
        keys = [q.key for q in qs]
        assert "kontraktvaerdi_bucket" in keys
        assert "anskaffelsesvej" in keys
        assert "fagomraade_saerlov" in keys
        assert "proces_status" in keys
        # Eksisterende stadig der
        assert "scope" in keys and "tilgang" in keys

    def test_apply_kontraktvaerdi_buckets(self):
        from src.services.risk_assessment.clarifying import apply_answers
        f1 = apply_answers(SystemFacts(), {"kontraktvaerdi_bucket": "Under 1,6 mio. kr."})
        assert f1.er_over_udbudsterskel() is False
        f2 = apply_answers(SystemFacts(), {"kontraktvaerdi_bucket": "1,6 - 5 mio. kr."})
        assert f2.er_over_udbudsterskel() is True
        f3 = apply_answers(SystemFacts(), {"kontraktvaerdi_bucket": "Over 5 mio. kr."})
        assert f3.er_over_udbudsterskel() is True
        f4 = apply_answers(SystemFacts(), {"kontraktvaerdi_bucket": "Ved ikke endnu"})
        assert f4.er_over_udbudsterskel() is None

    def test_apply_anskaffelsesvej_alle_varianter(self):
        from src.services.risk_assessment.clarifying import apply_answers
        cases = [
            ("SKI - direkte tildeling", Anskaffelsesvej.SKI_DIREKTE),
            ("SKI - mini-udbud", Anskaffelsesvej.SKI_MINIUDBUD),
            ("Under tærskel (ingen udbudspligt)", Anskaffelsesvej.UNDER_TAERSKEL),
            ("EU-udbud", Anskaffelsesvej.EU_UDBUD),
            ("Bygge- og anlægsprojekt", Anskaffelsesvej.BYGGE_ANLAEG),
        ]
        for answer, expected in cases:
            f = apply_answers(SystemFacts(), {"anskaffelsesvej": answer})
            assert f.anskaffelsesvej == expected, f"answer={answer!r}"

    def test_apply_fagomraade_separator(self):
        from src.services.risk_assessment.clarifying import apply_answers
        # Med separator → split i fagområde + særlov-liste
        f = apply_answers(SystemFacts(), {"fagomraade_saerlov": "Beskæftigelse — LAB §17a, forvaltningsloven"})
        assert f.fagomraade == "Beskæftigelse"
        assert "LAB §17a" in f.saerlovgivning
        assert "forvaltningsloven" in f.saerlovgivning

    def test_apply_fagomraade_uden_separator(self):
        from src.services.risk_assessment.clarifying import apply_answers
        f = apply_answers(SystemFacts(), {"fagomraade_saerlov": "Sundhed"})
        assert f.fagomraade == "Sundhed"
        assert f.saerlovgivning == []

    def test_proces_punkter_single_source_of_truth(self):
        # Options i spørgsmålet SKAL være identiske med PROCES_PUNKTER-labels,
        # og hver attr skal eksistere på SystemFacts — ellers divergerer UI og parsing.
        from src.services.risk_assessment.models import PROCES_PUNKTER
        from src.services.risk_assessment.clarifying import build_questions
        assert len(PROCES_PUNKTER) == 9
        f = SystemFacts()
        for attr, label in PROCES_PUNKTER:
            assert hasattr(f, attr), f"SystemFacts mangler {attr}"
            assert "," not in label, f"Label {label!r} indeholder komma — bryder split"
        qs = build_questions(SystemFacts(hosting_lokation="Azure"))
        proces_q = next(q for q in qs if q.key == "proces_status")
        assert proces_q.options == [label for _, label in PROCES_PUNKTER]

    def test_apply_kontraktvaerdi_saetter_estimat_flag(self):
        from src.services.risk_assessment.clarifying import apply_answers
        f = apply_answers(SystemFacts(), {"kontraktvaerdi_bucket": "Over 5 mio. kr."})
        assert f.kontraktvaerdi_er_estimat is True
        assert f.kontraktvaerdi_bucket_label == "Over 5 mio. kr."
        assert "Over 5 mio. kr." in f.kontraktvaerdi_label()
        assert "estimat" in f.kontraktvaerdi_label()

    def test_kontraktvaerdi_label_faktisk_beloeb(self):
        # Værdi fra dokument-ekstraktion (ikke bucket) → vis det faktiske tal
        f = SystemFacts(kontraktvaerdi_4aar_kr=2_400_000, kontraktvaerdi_er_estimat=False)
        label = f.kontraktvaerdi_label()
        assert "2.400.000" in label
        assert "estimat" not in label

    def test_ingen_kategori_konflikt_droppes(self):
        from src.services.risk_assessment.clarifying import apply_answers
        # "ingen" + konkrete kategorier → konkrete vinder
        f = apply_answers(SystemFacts(), {"persondata_kategorier": "ingen,cpr"})
        assert DataKategori.CPR in f.persondata_kategorier
        assert DataKategori.INGEN not in f.persondata_kategorier
        # "ingen" alene → bevares
        f2 = apply_answers(SystemFacts(), {"persondata_kategorier": "ingen"})
        assert f2.persondata_kategorier == [DataKategori.INGEN]

    def test_apply_proces_status_multiselect(self):
        from src.services.risk_assessment.clarifying import apply_answers
        f = apply_answers(SystemFacts(), {
            "proces_status": (
                "Digitalisering og IT adviseret tidligt,"
                "CIO har underskrevet kontrakt + DBA,"
                "DPIA-udkast sendt til DPO,"
                "AI-færdigheder dokumenteret (AI-forord. art. 4)"
            ),
        })
        assert f.dit_involveret_tidligt is True
        assert f.cio_har_underskrevet is True
        assert f.dpia_sendt_til_dpo is True
        assert f.ai_faerdigheder_dokumenteret is True
        # Ikke valgte forbliver False
        assert f.styregruppe_etableret is False
        assert f.contract_management_plan is False


# ---- document_extract ----------------------------------------------------


class TestDocumentExtract:
    def test_txt(self):
        d = extract_document("note.txt", b"hej verden")
        assert d.text == "hej verden"
        assert d.kind == "txt"
        assert d.char_count == 10

    def test_docx_roundtrip(self):
        # Byg et lille docx i hukommelsen
        doc = Document()
        doc.add_paragraph("Leverandør: Acme ApS")
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "CPR"
        t.rows[0].cells[1].text = "Ja"
        buf = io.BytesIO()
        doc.save(buf)
        d = extract_document("dba.docx", buf.getvalue())
        assert "Acme ApS" in d.text
        assert "CPR" in d.text  # tabel-indhold med

    def test_rejects_unsupported(self):
        with pytest.raises(DocumentExtractError):
            extract_document("image.png", b"\x89PNG")

    def test_rejects_oversized(self):
        big = b"x" * (11 * 1024 * 1024)
        with pytest.raises(DocumentExtractError):
            extract_document("big.txt", big)

    def test_rejects_empty(self):
        with pytest.raises(DocumentExtractError):
            extract_document("empty.txt", b"")

    def test_combine_truncates(self):
        from src.services.risk_assessment.document_extract import ExtractedDoc
        docs = [ExtractedDoc("a.txt", "x" * 50000, 50000, "txt")]
        combined = combine_for_prompt(docs, max_chars_per_doc=1000)
        assert "trunkeret" in combined
        assert "a.txt" in combined


# ---- docx_assembler + verifier (mod RIGTIG master-template) --------------


def _full_rv(systemnavn="Testsystem", n_risks=8):
    risks = [
        Risiko(
            risiko=f"Risiko {i}",
            konsekvens_beskrivelse=f"Konsekvens {i} for de registrerede.",
            sandsynlighed=Niveau.LAV_MIDDEL,
            score=Niveau.MIDDEL,
            hvorfor=f"Begrundelse {i}.",
        )
        for i in range(1, n_risks + 1)
    ]
    return Risikovurdering(
        facts=SystemFacts(systemnavn=systemnavn, leverandoer_navn="TestCorp"),
        risici=risks,
        formaal_tekst="Formålet er at vurdere systemet.",
        omfang_tekst="Vurderingen dækker hele systemet.",
        ansvarlige_tekst="Partheepan Vijayamohan og Anne Dandanell.",
        baggrund_tekst="Projektets baggrund er behovet for X.",
        funktionalitet_tekst="Systemet gør Y og Z.",
        interessenter_tekst="Dataansvarlig: Kalundborg Kommune.",
        personoplysninger_tekst="Almindelige personoplysninger og CPR.",
        lokationer_tekst="Microsoft Azure EU-region.",
        adgangsrettigheder_tekst="Entra ID SSO + MFA.",
        saarbarheder_tekst="Ingen kritiske sårbarheder.",
        tiltag_tekst="Databehandleraftale og kryptering.",
        ansvarlige_tiltag_tekst="Partheepan er ansvarlig.",
        kontrolmekanismer_tekst="Månedlig log-gennemgang.",
        opdatering_tekst="Revideres årligt.",
    )


@pytest.mark.skipif(
    not os.path.exists(get_template_path()),
    reason="master-template ikke til stede",
)
class TestAssembler:
    def test_assemble_produces_valid_docx(self):
        data = assemble_docx(_full_rv())
        assert data[:2] == b"PK"  # docx = zip
        assert len(data) > 10000

    def test_systemnavn_in_title_and_intro(self):
        data = assemble_docx(_full_rv(systemnavn="Voicecraft"))
        doc = Document(io.BytesIO(data))
        assert "Voicecraft" in doc.paragraphs[0].text
        assert "Voicecraft" in doc.paragraphs[3].text
        assert "Plan2learn" not in doc.paragraphs[3].text

    def test_risk_table_row_count(self):
        data = assemble_docx(_full_rv(n_risks=10))
        doc = Document(io.BytesIO(data))
        # header + 10 risici
        assert len(doc.tables[5].rows) == 11

    def test_no_bracket_artifacts_when_all_filled(self):
        data = assemble_docx(_full_rv())
        doc = Document(io.BytesIO(data))
        # Table 6 har legitim guidance "Eksempler [OBS: ...]" (yp[1] — rør ikke),
        # så den udelades fra streng bracket-check. Table 0,1,2,7 må være rene.
        for ti in [0, 1, 2, 7]:
            txt = doc.tables[ti].rows[0].cells[0].text
            assert "[" not in txt, f"Table {ti} har '[' artefakt: {txt[:80]}"
            assert "]" not in txt, f"Table {ti} har ']' artefakt: {txt[:80]}"
            assert ".." not in txt, f"Table {ti} har dobbelt-punktum: {txt[:80]}"
        # Table 6: verificér at de FYLDTE felter landede rent (uden brackets)
        t6 = doc.tables[6].rows[0].cells[0].text
        assert "Databehandleraftale og kryptering." in t6
        assert "Partheepan er ansvarlig." in t6

    def test_guidance_tables_untouched(self):
        data = assemble_docx(_full_rv())
        doc = Document(io.BytesIO(data))
        # Table 3 + 4 er "Eksempler"-guidance — skal stadig være der
        assert "Eksempler på risici" in doc.tables[3].rows[0].cells[0].text
        assert "Eksempler på konsekvenser" in doc.tables[4].rows[0].cells[0].text

    def test_verify_passes_for_complete_doc(self):
        data = assemble_docx(_full_rv(n_risks=8))
        res = verify_docx(data, expected_n_risks=8)
        assert res.valid, res.problems
        assert res.n_risk_rows == 8

    def test_verify_catches_wrong_risk_count(self):
        data = assemble_docx(_full_rv(n_risks=8))
        res = verify_docx(data, expected_n_risks=10)
        assert not res.valid
        assert any("forventet 10" in p for p in res.problems)

    def test_output_filename(self):
        assert output_filename("Velatir") == "Databeskyttelsesretlig risikovurdering - Velatir.docx"
        assert "/" not in output_filename("A/B")

    def test_output_filename_sanitizes_windows_forbidden(self):
        # Windows afviser \ / : * ? " < > | i filnavne
        navn = output_filename('Sys/V:1*x?"<y>|z\\w')
        for ch in '\\/:*?"<>|':
            assert ch not in navn.replace(".docx", "").split(" - ")[1], f"{ch!r} ikke sanitized"
        # Tomt navn → fallback
        assert output_filename("") == "Databeskyttelsesretlig risikovurdering - system.docx"
        assert output_filename("***") == "Databeskyttelsesretlig risikovurdering - system.docx"


@pytest.mark.skipif(
    not os.path.exists(get_template_path()),
    reason="master-template ikke til stede",
)
class TestVerifierKalundborgCompliance:
    """Etape 3: verifier flagger kommunale compliance-issues når facts gives."""

    def _rv_med_facts(self, **fact_kwargs):
        rv = _full_rv()
        for k, v in fact_kwargs.items():
            setattr(rv.facts, k, v)
        return rv

    def test_uden_facts_kører_kun_basis_check(self):
        # Bagudkompatibelt: ingen facts → ingen kommunale checks
        data = assemble_docx(_full_rv())
        res = verify_docx(data, expected_n_risks=8)
        assert res.valid

    def test_udbudspligt_mismatch_flagges_hvis_ikke_naevnt(self):
        # Fakta: over tærskel + SKI (mismatch). Standard tiltag-tekst nævner ikke EU-udbud.
        rv = self._rv_med_facts(
            kontraktvaerdi_4aar_kr=3_000_000,
            anskaffelsesvej=Anskaffelsesvej.SKI_DIREKTE,
        )
        # Sørg for at teksten IKKE nævner EU-udbud
        rv.tiltag_tekst = "Generel kryptering og adgangskontrol."
        rv.ansvarlige_tekst = "Partheepan Vijayamohan, AI Program Lead."
        # ansvarlige_tekst skal nævne kommunal aktør for at undgå check 6b
        rv.ansvarlige_tekst = "DPO inddrages. Partheepan Vijayamohan."
        rv.saarbarheder_tekst = "Ingen kendte."
        data = assemble_docx(rv)
        res = verify_docx(data, expected_n_risks=8, facts=rv.facts)
        assert not res.valid
        assert any("EU-udbud" in p or "udbudspligt" in p for p in res.problems)

    def test_kommunale_aktorer_mangler_flagges(self):
        rv = self._rv_med_facts()
        # Strip kommunale aktører helt fra ansvarlige
        rv.ansvarlige_tekst = "Partheepan og Anne er ansvarlige."
        data = assemble_docx(rv)
        res = verify_docx(data, expected_n_risks=8, facts=rv.facts)
        assert any("kommunale aktører" in p or "kommunal kontekst" in p for p in res.problems)

    def test_kommunale_aktorer_med_cio_passerer(self):
        rv = self._rv_med_facts()
        rv.ansvarlige_tekst = "Digitaliserings- og IT-chefen underskriver kontrakten."
        data = assemble_docx(rv)
        res = verify_docx(data, expected_n_risks=8, facts=rv.facts)
        # Ingen "kommunale aktører"-problem
        assert not any("kommunale aktører" in p for p in res.problems)

    def test_fagomraade_uden_saerlov_naevnt_flagges(self):
        rv = self._rv_med_facts(fagomraade="Beskæftigelse", saerlovgivning=["LAB §17a"])
        # Sørg for at INGEN tekst nævner fagområdet eller særlov
        for attr in ("formaal_tekst", "omfang_tekst", "ansvarlige_tekst", "baggrund_tekst",
                     "funktionalitet_tekst", "interessenter_tekst", "personoplysninger_tekst",
                     "lokationer_tekst", "adgangsrettigheder_tekst", "saarbarheder_tekst",
                     "tiltag_tekst", "ansvarlige_tiltag_tekst", "kontrolmekanismer_tekst",
                     "opdatering_tekst"):
            v = getattr(rv, attr)
            # Behold en kommunal aktør i ansvarlige så check 6b passerer
            if attr == "ansvarlige_tekst":
                setattr(rv, attr, "DPO inddrages.")
            else:
                setattr(rv, attr, v.replace("Beskæftigelse", "X").replace("LAB", "X").replace("særlov", "X"))
        data = assemble_docx(rv)
        res = verify_docx(data, expected_n_risks=8, facts=rv.facts)
        assert any("fagområde" in p.lower() or "fagomraade" in p.lower() or "særlov" in p for p in res.problems)
